#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Walidator macierzy funkcjonalnosci (wspolny schemat iMap) z obsluga profili produktow.
Uzycie: python3 sprawdz_macierz.py <plik.xlsx|plik.csv> [--arkusz Macierz] [--rejestr rejestr.csv] [--profil profile/<produkt>.json]
Bez --profil dziala jak dotad (wspolny schemat). NIE modyfikuje danych, niczego nie zmysla."""
import sys, os, json, csv, re, argparse

BASE_HEADERS = ["Lp","ID","Poziom 1 – Obszar","Poziom 2 – Grupa","Funkcjonalność","Opis","Opis biznesowy",
"Typ","Interfejs","Wersja / standard","Identyfikator techniczny","ID funkcji","Dokumentacja","Źródło","Status","Rodzaj","Projekt","Uwagi"]
BASE_KEYS = ["lp","id_produktu","poziom1","poziom2","funkcjonalnosc","opis","opis_biznesowy","typ","interfejs",
"wersja_standard","identyfikator_techniczny","id_funkcji","dokumentacja","zrodlo","status","rodzaj","projekt","uwagi"]
CORE_REQUIRED = ["poziom1","poziom2","funkcjonalnosc","typ","zrodlo","status"]
TYP_DEFAULT = {"widget","tool","action","złożona","usługa","operacja API","standard OGC","zadanie GP",
"funkcja silnika","format","źródło danych","mechanizm bezpieczeństwa","funkcja administracyjna","moduł","wdrożenie"}
IFACE = {"REST","SOAP","WMS","WMTS","WFS","WPS","OGC API Features","—"}
DOK = {"TAK","NIE"}; ZR = {"kod","konfiguracja","dokumentacja","baza wiedzy"}; ST = {"potwierdzone","do weryfikacji"}
RODZAJ = {"podstawowa","rozszerzona","dedykowana"}
PROJEKT = {"MIIP","Centra Operacyjne","RIIP","MineSens","GEOINFO2","ORLEN","GDOS","WIR","HydroIQ","GUGIK","NSIS",
"ARiMR/NGIS","UM Nowa Sól","ISOK","MZP/MRP","GAZ","CAPAP","EMUiA","ORLEN VIA","OZE","UMWL","UM Świdnica",
"UM Wodzisław Śląski","Wojsko","AMS","UM Racibórz","UM Kołobrzeg","GEOINFO3","WROSIP","SIPWW","GUS","iMapCloud SaaS"}
ID_FUNC_RE = re.compile(r"^FN-\d{4,}$")

def load_profile(path):
    if not path: return {}
    return json.load(open(path, encoding="utf-8"))

def build_schema(profile):
    h2k = {h:k for h,k in zip(BASE_HEADERS, BASE_KEYS)}
    for h in profile.get("kolumny_wylaczone", []): h2k.pop(h, None)
    for h,k in profile.get("aliasy_kolumn", {}).items(): h2k[h]=k
    extra_dicts={}
    for col in profile.get("kolumny_dodatkowe", []):
        h2k[col["naglowek"]]=col["klucz"]
        if col.get("slownik"): extra_dicts[col["klucz"]]=set(col["slownik"])
    typ_set = set(profile["slownik_typ"]) if profile.get("slownik_typ") else set(TYP_DEFAULT)
    known_keys = set(BASE_KEYS) | {c["klucz"] for c in profile.get("kolumny_dodatkowe", [])}
    return {"h2k":h2k,"typ":typ_set,"extra":extra_dicts,"prefix":profile.get("prefiks"),
            "zrodlo_free":bool(profile.get("zrodlo_wolny_tekst")),"known_keys":known_keys,
            "known_headers":set(h2k.keys())}

def hdr2key(h, sch):
    if h in PROJEKT: return "PRJ:"+h
    return h if h in sch["known_keys"] else sch["h2k"].get(h)

def load(path, sheet, sch):
    ext=os.path.splitext(path)[1].lower(); rows=[]
    if ext==".csv":
        with open(path, encoding="utf-8-sig") as f:
            rd=csv.DictReader(f); header=rd.fieldnames or []
            for d in rd:
                row={}
                for h in header:
                    k=hdr2key(h, sch)
                    if k: row[k]=(d.get(h,"") or "").strip()
                rows.append(row)
        return rows, header
    from openpyxl import load_workbook
    wb=load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else (wb["Macierz"] if "Macierz" in wb.sheetnames else wb.active)
    data=list(ws.iter_rows(values_only=True)); hidx=None; header=None
    for i,r in enumerate(data):
        vals=[("" if v is None else str(v)).strip() for v in r]
        mapped={hdr2key(h, sch) for h in vals}
        if "funkcjonalnosc" in mapped and "poziom1" in mapped:
            hidx=i; header=vals; break
    if hidx is None:
        return [], "Nie znaleziono naglowka zgodnego ze schematem/profilem."
    colmap={}
    for j,h in enumerate(header):
        k=hdr2key(h, sch)
        if k and k not in colmap: colmap[k]=j
    for r in data[hidx+1:]:
        vals=[("" if v is None else str(v)).strip() for v in r]
        if not any(vals): continue
        rows.append({k:(vals[j] if j<len(vals) else "") for k,j in colmap.items()})
    return rows, header

def load_registry(path):
    ids=set()
    if not path: return ids
    with open(path, encoding="utf-8") as f:
        for d in csv.DictReader(f):
            v=(d.get("id_funkcji","") or "").strip()
            if v: ids.add(v)
    return ids

def validate(rows, header, sch, registry=None):
    errors=[]; warnings=[]; registry=registry or set()
    typ_set=sch["typ"]; extra=sch["extra"]; prefix=sch["prefix"]
    iface_active = any(hdr2key(h,sch)=="interfejs" for h in (header or []))
    id_re = re.compile("^%s-(\\d+\\.\\d+\\.\\d{2,}|\\d{2}-\\d{2}-\\d{4}|\\d{3,})$" % re.escape(prefix)) if prefix else re.compile(r"^[A-Z]{2,4}-(\d+\.\d+\.\d{2,}|\d{2}-\d{2}-\d{4}|\d{3,})$")
    if header is not None and not isinstance(header,str):
        covered={hdr2key(h,sch) for h in header}
        miss=[k for k in CORE_REQUIRED if k not in covered]
        if miss: errors.append("Brak wymaganych kolumn (klucze): %s" % miss)
        extra_h=[h for h in header if h and hdr2key(h,sch) is None]
        if extra_h: warnings.append("Kolumny spoza schematu/profilu (zignorowane): %s" % extra_h)
    seen={}; seen_idp={}
    for i,r in enumerate(rows, start=1):
        n=i; weryf=r.get("status","")=="do weryfikacji"
        for k in CORE_REQUIRED:
            if not r.get(k):
                if weryf and k=="typ": continue
                errors.append("Wiersz %d: puste wymagane pole '%s'." % (n,k))
        if r.get("typ") and r["typ"] not in typ_set:
            errors.append("Wiersz %d: Typ '%s' spoza slownika." % (n,r["typ"]))
        if iface_active:
            for tok in [t.strip() for t in (r.get("interfejs","") or "").replace("/",";").split(";") if t.strip()]:
                if tok not in IFACE: errors.append("Wiersz %d: Interfejs '%s' spoza slownika." % (n,tok))
        if r.get("dokumentacja") and r["dokumentacja"] not in DOK:
            errors.append("Wiersz %d: Dokumentacja '%s' spoza TAK/NIE." % (n,r["dokumentacja"]))
        if r.get("status") and r["status"] not in ST:
            errors.append("Wiersz %d: Status '%s' spoza slownika." % (n,r["status"]))
        if r.get("rodzaj") and r["rodzaj"] not in RODZAJ:
            errors.append("Wiersz %d: Rodzaj '%s' spoza slownika (podstawowa/rozszerzona/dedykowana)." % (n,r["rodzaj"]))
        for tok in [t.strip() for t in (r.get("projekt","") or "").split(";") if t.strip()]:
            if tok not in PROJEKT:
                warnings.append("Wiersz %d: Projekt '%s' spoza listy (dopisz do Slownikow, jesli nowy)." % (n,tok))
        for pk,pv in r.items():
            if pk.startswith("PRJ:") and pv and pv not in ("✓","x","X","TAK","1"):
                warnings.append("Wiersz %d: projekt '%s' ma nietypowa wartosc '%s' (oczekiwano ✓ lub pusto)." % (n,pk[4:],pv))
        if not sch["zrodlo_free"]:
            for tok in [t.strip() for t in (r.get("zrodlo","") or "").split(";") if t.strip()]:
                if tok.split("(")[0].strip() not in ZR:
                    warnings.append("Wiersz %d: Zrodlo '%s' spoza slownika." % (n,tok))
        for k,dset in extra.items():
            v=r.get(k,"")
            if v and v not in dset:
                errors.append("Wiersz %d: kolumna '%s' wartosc '%s' spoza slownika." % (n,k,v))
        idp=r.get("id_produktu","")
        if idp:
            if not id_re.match(idp):
                errors.append("Wiersz %d: ID '%s' niezgodne z wymaganym formatem." % (n,idp))
            if idp in seen_idp: errors.append("Wiersz %d: zduplikowane ID '%s' (juz w wierszu %d)." % (n,idp,seen_idp[idp]))
            else: seen_idp[idp]=n
        idf=r.get("id_funkcji","")
        if idf:
            if not ID_FUNC_RE.match(idf): errors.append("Wiersz %d: ID funkcji '%s' niezgodne z FN-NNNNN." % (n,idf))
            elif registry and idf not in registry: errors.append("Wiersz %d: ID funkcji '%s' nie istnieje w rejestrze." % (n,idf))
        if r.get("status")=="potwierdzone":
            if not r.get("zrodlo"): warnings.append("Wiersz %d: potwierdzone bez Zrodla." % n)
            if not (r.get("identyfikator_techniczny") or r.get("opis") or r.get("opis_biznesowy")):
                warnings.append("Wiersz %d: potwierdzone bez zakotwiczenia (id/opis)." % n)
        if iface_active and r.get("wersja_standard") and not r.get("interfejs"):
            warnings.append("Wiersz %d: Wersja/standard bez Interfejsu." % n)
        key=(r.get("funkcjonalnosc",""), r.get("typ",""), r.get("identyfikator_techniczny",""))
        if key in seen and any(key): warnings.append("Wiersz %d: mozliwy duplikat wiersza %d." % (n,seen[key]))
        else: seen[key]=n
    from collections import Counter
    stats={"wiersze":len(rows),
           "z_id_produktu":sum(1 for r in rows if r.get("id_produktu")),
           "z_id_funkcji":sum(1 for r in rows if r.get("id_funkcji")),
           "do_weryfikacji":sum(1 for r in rows if r.get("status")=="do weryfikacji"),
           "dokumentacja_NIE":sum(1 for r in rows if r.get("dokumentacja")=="NIE"),
           "wg_poziom1":dict(Counter(r.get("poziom1","") for r in rows)),
           "wg_typ":dict(Counter(r.get("typ","") for r in rows))}
    return errors, warnings, stats

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("plik"); ap.add_argument("--arkusz",default="Macierz")
    ap.add_argument("--rejestr",default=None); ap.add_argument("--profil",default=None)
    a=ap.parse_args()
    sch=build_schema(load_profile(a.profil))
    rows,header=load(a.plik,a.arkusz,sch)
    if isinstance(header,str) and not rows:
        print(json.dumps({"status":"errors_found","errors":[header]},ensure_ascii=False,indent=2)); sys.exit(1)
    errors,warnings,stats=validate(rows,header,sch,load_registry(a.rejestr))
    out={"status":"ok" if not errors else "errors_found","plik":os.path.basename(a.plik),
         "profil":os.path.basename(a.profil) if a.profil else None,
         "liczba_bledow":len(errors),"liczba_ostrzezen":len(warnings),
         "errors":errors[:200],"warnings":warnings[:200],"stats":stats}
    print(json.dumps(out,ensure_ascii=False,indent=2)); print("")
    print("=== RAPORT ===")
    print("Plik: %s | profil: %s | wiersze: %d | bledy: %d | ostrzezenia: %d" % (
        out["plik"], out["profil"] or "(wspolny)", stats["wiersze"], len(errors), len(warnings)))
    print("Werdykt: " + ("GOTOWA (brak bledow)" if not errors else "DO POPRAWY (sa bledy)"))

if __name__=="__main__": main()
