#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Porownanie dwoch macierzy funkcjonalnosci (stara vs nowa).
Uzycie: python3 porownaj_macierze.py --stara <plik> --nowa <plik> [--klucz auto|nazwa|idfunkcji]
Zwraca: funkcje nowe / usuniete / wspolne / zmienione (+ JSON). Nie modyfikuje danych."""
import os, csv, json, argparse

KEYS=["lp","id_produktu","poziom1","poziom2","funkcjonalnosc","opis","opis_biznesowy","typ","interfejs",
"wersja_standard","identyfikator_techniczny","id_funkcji","dokumentacja","zrodlo","status","uwagi"]
HEADERS=["Lp","ID","Poziom 1 – Obszar","Poziom 2 – Grupa","Funkcjonalność","Opis","Opis biznesowy","Typ",
"Interfejs","Wersja / standard","Identyfikator techniczny","ID funkcji","Dokumentacja","Źródło","Status","Uwagi"]
POROWN=["poziom1","poziom2","opis","opis_biznesowy","typ","identyfikator_techniczny"]

def load(path):
    ext=os.path.splitext(path)[1].lower(); rows=[]
    if ext==".csv":
        with open(path,encoding="utf-8") as f:
            for d in csv.DictReader(f):
                rows.append({k:(d.get(k,"") or "").strip() for k in KEYS})
        return rows
    from openpyxl import load_workbook
    wb=load_workbook(path, read_only=True, data_only=True)
    ws=wb["Macierz"] if "Macierz" in wb.sheetnames else wb.active
    data=list(ws.iter_rows(values_only=True)); hidx=None; header=None
    for i,r in enumerate(data):
        vals=[("" if v is None else str(v)).strip() for v in r]
        if "Funkcjonalność" in vals and "Poziom 1 – Obszar" in vals:
            hidx=i; header=vals; break
    if hidx is None: return rows
    cm={h:j for j,h in enumerate(header)}
    for r in data[hidx+1:]:
        vals=[("" if v is None else str(v)).strip() for v in r]
        if not any(vals): continue
        row={}
        for k,h in zip(KEYS,HEADERS):
            j=cm.get(h); row[k]=vals[j] if (j is not None and j<len(vals)) else ""
        rows.append(row)
    return rows

def norm(s): return " ".join((s or "").lower().split())

def keyfn(mode, rows_a, rows_b):
    if mode=="idfunkcji": return lambda r: r.get("id_funkcji","")
    if mode=="nazwa": return lambda r: norm(r.get("funkcjonalnosc",""))+"|"+r.get("typ","")
    # auto: id_funkcji jesli obie strony maja komplet, inaczej nazwa+typ
    def has_idf(rows): return rows and all(r.get("id_funkcji") for r in rows)
    if has_idf(rows_a) and has_idf(rows_b):
        return lambda r: r.get("id_funkcji","")
    return lambda r: norm(r.get("funkcjonalnosc",""))+"|"+r.get("typ","")

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--stara",required=True); ap.add_argument("--nowa",required=True)
    ap.add_argument("--klucz",default="auto",choices=["auto","nazwa","idfunkcji"])
    a=ap.parse_args()
    A=load(a.stara); B=load(a.nowa)
    kf=keyfn(a.klucz,A,B)
    da={kf(r):r for r in A if kf(r)}; db={kf(r):r for r in B if kf(r)}
    nowe=[db[k] for k in db if k not in da]
    usuniete=[da[k] for k in da if k not in db]
    wspolne=[k for k in db if k in da]
    zmienione=[]
    for k in wspolne:
        diff={f:(da[k].get(f,""),db[k].get(f,"")) for f in POROWN if (da[k].get(f,"") or "")!=(db[k].get(f,"") or "")}
        if diff: zmienione.append({"funkcja":db[k].get("funkcjonalnosc",""),"zmiany":diff})
    out={"stara":os.path.basename(a.stara),"nowa":os.path.basename(a.nowa),"klucz":a.klucz,
         "liczba_stara":len(A),"liczba_nowa":len(B),
         "nowe":len(nowe),"usuniete":len(usuniete),"wspolne":len(wspolne),"zmienione":len(zmienione),
         "nowe_nazwy":[r.get("funkcjonalnosc","") for r in nowe][:50],
         "usuniete_nazwy":[r.get("funkcjonalnosc","") for r in usuniete][:50],
         "zmienione_szczegoly":zmienione[:50]}
    print(json.dumps(out,ensure_ascii=False,indent=2))
    print("")
    print("=== POROWNANIE ===")
    print("stara=%d | nowa=%d | nowe=%d | usuniete=%d | wspolne=%d | zmienione=%d" % (
        len(A),len(B),len(nowe),len(usuniete),len(wspolne),len(zmienione)))

if __name__=="__main__":
    main()
